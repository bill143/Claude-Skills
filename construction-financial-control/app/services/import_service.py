"""Bulk budget import from Excel (.xlsx) or CSV.

Flow: upload -> parse -> flexible column mapping (header aliases) -> normalize
CSI codes -> validate against the cost-code library -> preview or commit.
Preview never writes; commit creates budget lines in one transaction and
journals the import in the audit ledger.
"""
import csv
import io
from decimal import Decimal, InvalidOperation

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.models.budget import BudgetLine
from app.models.enums import CostCategory
from app.models.project import Project
from app.models.user import User
from app.services import audit_service, wbs_service
from app.services.forecast_service import code_digits, division_of, normalize_code

CODE_ALIASES = {"code", "cost code", "cost_code", "costcode", "csi", "csi code", "wbs",
                "wbs code", "item", "section", "04 section"}
DESCRIPTION_ALIASES = {"description", "title", "name", "scope", "item description", "04 title"}
AMOUNT_ALIASES = {"amount", "budget", "original budget", "original_budget", "value",
                  "total", "cost", "budget amount", "estimate"}
CATEGORY_ALIASES = {"category", "type", "cost type"}

MAX_ROWS = 5000


def _clean_header(value) -> str:
    return str(value).strip().lower().replace("\xa0", " ") if value is not None else ""


def _parse_amount(value) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    text = str(value).strip().replace("$", "").replace(",", "").replace("\xa0", "")
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    try:
        amount = Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None
    return -amount if negative else amount


def _rows_from_file(filename: str, content: bytes) -> list[list]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            import openpyxl
        except ImportError as exc:  # pragma: no cover
            raise HTTPException(status_code=500, detail="openpyxl not installed") from exc
        workbook = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        worksheet = workbook[workbook.sheetnames[0]]
        return [list(row) for row in worksheet.iter_rows(values_only=True)]
    if name.endswith((".csv", ".txt")) or not name:
        text = content.decode("utf-8-sig", errors="replace")
        return [row for row in csv.reader(io.StringIO(text))]
    raise HTTPException(status_code=422, detail="Unsupported file type: upload .xlsx or .csv")


def _find_columns(rows: list[list]) -> tuple[int, dict]:
    """Locate the header row and map column indexes by alias."""
    for row_index, row in enumerate(rows[:20]):
        headers = [_clean_header(cell) for cell in row]
        mapping = {}
        for col_index, header in enumerate(headers):
            if header in CODE_ALIASES and "code" not in mapping:
                mapping["code"] = col_index
            elif header in DESCRIPTION_ALIASES and "description" not in mapping:
                mapping["description"] = col_index
            elif header in AMOUNT_ALIASES and "amount" not in mapping:
                mapping["amount"] = col_index
            elif header in CATEGORY_ALIASES and "category" not in mapping:
                mapping["category"] = col_index
        if "code" in mapping and "amount" in mapping:
            return row_index, mapping
    raise HTTPException(
        status_code=422,
        detail="Could not find header row: need a cost-code column "
               f"({sorted(CODE_ALIASES)[:4]}...) and an amount column "
               f"({sorted(AMOUNT_ALIASES)[:4]}...)",
    )


def _infer_category(explicit, division: str) -> CostCategory:
    if explicit:
        cleaned = str(explicit).strip().upper().replace(" ", "_").replace("/", "_")
        try:
            return CostCategory(cleaned)
        except ValueError:
            pass
    return CostCategory.GENERAL_CONDITIONS if division in ("00", "01") else CostCategory.SUBCONTRACT


def analyze(db: Session, project: Project, filename: str, content: bytes) -> dict:
    """Parse + validate the file; returns row dispositions without writing."""
    rows = _rows_from_file(filename, content)
    if len(rows) > MAX_ROWS + 20:
        raise HTTPException(status_code=422, detail=f"File exceeds {MAX_ROWS} rows")
    header_index, columns = _find_columns(rows)

    existing_codes = {
        code_digits(code) for (code,) in
        db.execute(select(BudgetLine.cost_code).where(BudgetLine.project_id == project.id))
    }

    def cell(row, key):
        index = columns.get(key)
        return row[index] if index is not None and index < len(row) else None

    parsed: dict[str, dict] = {}
    invalid: list[dict] = []
    for offset, row in enumerate(rows[header_index + 1:], start=header_index + 2):
        raw_code = cell(row, "code")
        if raw_code is None or not str(raw_code).strip():
            continue  # blank/spacer row
        digits = code_digits(raw_code)
        amount = _parse_amount(cell(row, "amount"))
        problem = None
        if len(digits) < 2:
            problem = "cost code needs at least a two-digit division"
        elif amount is None:
            problem = "amount is missing or not a number"
        elif amount < 0:
            problem = "amount cannot be negative"
        if problem:
            invalid.append({"row": offset, "code": str(raw_code), "error": problem})
            continue
        assert amount is not None  # narrowed by the checks above

        code = normalize_code(raw_code)
        division = division_of(code)
        description = cell(row, "description")
        description = (str(description).strip() if description else None) \
            or wbs_service.lookup_title(db, code) or f"Cost code {code}"
        entry = parsed.get(digits)
        if entry:  # same code twice in the file -> merge amounts
            entry["amount"] = float(Decimal(str(entry["amount"])) + amount)
            entry["merged_rows"] += 1
        else:
            parsed[digits] = {
                "code": code,
                "division": division,
                "description": description[:255],
                "category": _infer_category(cell(row, "category"), division).value,
                "amount": float(amount),
                "in_library": wbs_service.lookup_title(db, code) is not None,
                "exists_in_project": digits in existing_codes,
                "merged_rows": 1,
            }

    new_rows = [r for r in parsed.values() if not r["exists_in_project"]]
    duplicates = [r for r in parsed.values() if r["exists_in_project"]]
    return {
        "filename": filename,
        "rows_parsed": len(parsed),
        "rows_new": len(new_rows),
        "rows_duplicate": len(duplicates),
        "rows_invalid": len(invalid),
        "total_amount_new": round(sum(r["amount"] for r in new_rows), 2),
        "rows": sorted(parsed.values(), key=lambda r: r["code"]),
        "invalid": invalid[:50],
    }


def commit(db: Session, project: Project, filename: str, content: bytes, *,
           actor: User, on_duplicate: str = "skip") -> dict:
    """Create budget lines from the file. on_duplicate: skip | update | error."""
    if on_duplicate not in ("skip", "update", "error"):
        raise HTTPException(status_code=422, detail="on_duplicate must be skip, update, or error")
    preview = analyze(db, project, filename, content)
    if preview["rows_invalid"]:
        raise HTTPException(
            status_code=422,
            detail={"message": "Fix invalid rows before committing", "invalid": preview["invalid"]},
        )
    if on_duplicate == "error" and preview["rows_duplicate"]:
        raise HTTPException(
            status_code=422,
            detail={"message": "Duplicate cost codes already exist in this project",
                    "duplicates": [r["code"] for r in preview["rows"] if r["exists_in_project"]]},
        )

    existing = {
        code_digits(line.cost_code): line for line in
        db.execute(select(BudgetLine).where(BudgetLine.project_id == project.id)).scalars()
    }
    created = updated = skipped = 0
    for row in preview["rows"]:
        digits = code_digits(row["code"])
        line = existing.get(digits)
        if line is not None:
            if on_duplicate == "update":
                line.original_budget = Decimal(str(row["amount"]))
                line.description = row["description"]
                updated += 1
            else:
                skipped += 1
            continue
        db.add(BudgetLine(
            project_id=project.id,
            cost_code=row["code"],
            description=row["description"],
            category=CostCategory(row["category"]),
            original_budget=Decimal(str(row["amount"])),
        ))
        created += 1
    db.flush()
    audit_service.record(
        db, actor=actor, entity_type="PROJECT", entity_id=project.id,
        action="budget_imported",
        payload={"filename": filename, "created": created, "updated": updated,
                 "skipped": skipped, "total_amount": preview["total_amount_new"]},
    )
    db.commit()
    return {"created": created, "updated": updated, "skipped": skipped,
            "total_amount_new": preview["total_amount_new"]}
